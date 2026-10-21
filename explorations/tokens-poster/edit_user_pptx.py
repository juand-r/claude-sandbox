"""Add Method sections 4-5 and a Setup box to the user's Keynote export (babylm2026.pptx, Oct 21).

What it does, on a copy:
  * Method box (a hand-drawn frame picture) grows by GROW pt; the Results box starts GROW pt lower.
    Frames are resized by 9-slice scaling (corners and edges keep their thickness).
  * Method row 4 "What curation changed" (figures/curation.png) and 5 "Quality checks" (figures/quality.png),
    with number circles and headings copied from section 3.
  * Results contents move down; the chart is swapped for figures/results-wide.png; a Setup box
    (white box + frame, like "What we did") goes under the two result paragraphs.

Usage: python edit_user_pptx.py IN.pptx OUT.pptx
"""
import copy
import io
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Pt

HERE = Path(__file__).parent
GROW = 220                      # pt added to the Method box and taken from the Results box
SLICE = 40                      # px kept unscaled at each frame edge (frame images are 1 px per pt)
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def emu(pt):
    return int(round(pt * 12700))


def nine_slice(blob, new_w, new_h):
    """Resize a frame image so the SLICE-px border bands keep their size."""
    im = Image.open(io.BytesIO(blob)).convert("RGBA")
    w, h, s = im.width, im.height, SLICE
    out = Image.new("RGBA", (new_w, new_h))
    xs = [(0, s, 0, s), (s, w - s, s, new_w - s), (w - s, w, new_w - s, new_w)]
    ys = [(0, s, 0, s), (s, h - s, s, new_h - s), (h - s, h, new_h - s, new_h)]
    for sx0, sx1, dx0, dx1 in xs:
        for sy0, sy1, dy0, dy1 in ys:
            tile = im.crop((sx0, sy0, sx1, sy1)).resize((dx1 - dx0, dy1 - dy0), Image.LANCZOS)
            out.paste(tile, (dx0, dy0))
    buf = io.BytesIO()
    out.save(buf, "PNG")
    buf.seek(0)
    return buf


class Slide:
    def __init__(self, slide):
        self.slide = slide
        self.tree = slide.shapes._spTree
        self.next_id = 1 + max(int(e.get("id")) for e in self.tree.iter(P + "cNvPr"))

    def get(self, name):
        found = [s for s in self.slide.shapes if s.name == name]
        if len(found) != 1:
            raise ValueError(f"expected one shape named {name!r}, found {len(found)}")
        return found[0]

    def move(self, name, dy, x=None, w=None, h=None):
        sh = self.get(name)
        sh.top = sh.top + emu(dy)
        if x is not None:
            sh.left = emu(x)
        if w is not None:
            sh.width = emu(w)
        if h is not None:
            sh.height = emu(h)

    def _fresh_ids(self, el):
        for c in el.iter(P + "cNvPr"):
            c.set("id", str(self.next_id))
            self.next_id += 1

    def clone(self, name, new_name, x, y, text=None):
        """Copy a shape or group, place it at (x, y) pt, optionally replace its text."""
        src = self.get(name)._element
        el = copy.deepcopy(src)
        self._fresh_ids(el)
        el.find(".//" + P + "cNvPr").set("name", new_name)
        xfrm = el.find(".//" + A + "xfrm")
        xfrm.find(A + "off").set("x", str(emu(x)))
        xfrm.find(A + "off").set("y", str(emu(y)))
        if text is not None:
            runs = el.findall(".//" + A + "r")
            runs[0].find(A + "t").text = text
            for r in runs[1:]:
                r.getparent().remove(r)
        self.tree.append(el)
        return el

    def picture(self, path_or_buf, name, x, y, w, h, replace=None):
        old = self.get(replace)._element if replace is not None else None
        pic = self.slide.shapes.add_picture(path_or_buf, emu(x), emu(y), emu(w), emu(h))
        pic.name = name
        if old is not None:   # keep the old shape's stacking position
            old.addprevious(pic._element)
            old.getparent().remove(old)
        return pic


def main(src, dst):
    prs = Presentation(src)
    s = Slide(prs.slides[0])

    # 1. Frames: Method grows, Results starts lower and shrinks.
    for name, dy, dh in (("method-box", 0, GROW), ("results-box", GROW, -GROW)):
        sh = s.get(name)
        x, y, w, h = sh.left / 12700, sh.top / 12700, sh.width / 12700, sh.height / 12700
        im = Image.open(io.BytesIO(sh.image.blob))
        if abs(im.width - w) > 1.5 or abs(im.height - h) > 1.5:
            raise ValueError(f"{name}: image {im.size} is not about 1 px per pt ({w:.1f} x {h:.1f})")
        s.picture(nine_slice(sh.image.blob, im.width, round(im.height + dh)), name, x, y + dy, w, h + dh, replace=name)

    # 2. Method row: sections 4 and 5.
    row = 2505
    for num, x, title in (("4", 829, "What curation changed"), ("5", 1624, "Quality checks")):
        s.clone("m3-heading-number", f"m{num}-heading-number", x, row, text=num)
        head = s.clone("m3-heading", f"m{num}-heading", x + 56, row + 2, text=title)
        head.find(".//" + A + "xfrm").find(A + "ext").set("cx", str(emu(660)))   # one line
    s.clone("m1-text", "m4-text", 829, row + 51, text="28,571 explicit decisions; about 10,000 change MorphyNet:")
    s.picture(str(HERE / "figures/curation.png"), "curation", 829, row + 95, 723.75, 150)
    s.picture(str(HERE / "figures/quality.png"), "quality", 1624, row + 51, 723.75, 196.5)

    # 3. Results: move down, wide chart, Setup box.
    s.move("results-heading", GROW)
    s.move("leaderboard-text", GROW)
    s.move("curation-text", GROW - 15)
    s.move("chart-title", GROW)
    chart = s.get("results-chart")
    # chart-title ends at 2938 pt after the move; chart 2946-3225; note 3233-3286 (box ends 3304)
    s.picture(str(HERE / "figures/results-wide.png"), "results-chart", chart.left / 12700, 2946, 665, 279.3,
              replace="results-chart")
    note = s.get("chart-note")
    s.move("chart-note", 3233 - note.top / 12700)

    what = s.get("what-box")
    frame = [c for c in what.shapes if c.shape_type == 13][0]
    box_x, box_y, box_w, box_h = 829, 3146, 724, 138
    fill = copy.deepcopy([c for c in what.shapes if c.shape_type == 1][0]._element)
    s._fresh_ids(fill)
    fill.find(".//" + P + "cNvPr").set("name", "setup-box")
    fx = fill.find(".//" + A + "xfrm")
    fx.find(A + "off").set("x", str(emu(box_x + 2.5))); fx.find(A + "off").set("y", str(emu(box_y + 2.5)))
    fx.find(A + "ext").set("cx", str(emu(box_w - 5))); fx.find(A + "ext").set("cy", str(emu(box_h - 5)))
    s.tree.append(fill)
    fim = Image.open(io.BytesIO(frame.image.blob))
    scale = fim.width / (frame.width / 12700)          # frame image px per pt
    s.picture(nine_slice(frame.image.blob, round(box_w * scale), round(box_h * scale)), "setup-box", box_x, box_y, box_w, box_h)
    setup = s.clone("o1-note", "setup-text", box_x + 18, box_y + 14,
                    text="Llama-style decoder (12 layers, 768 hidden), trained for 10 epochs on the 10M-word "
                         "strict-small corpus with the Muon optimizer. Comparisons change only the tokenizer. "
                         "The submitted model (about 135M parameters) adds masked next-token prediction.")
    ext = setup.find(".//" + A + "xfrm").find(A + "ext")
    ext.set("cx", str(emu(box_w - 36)))
    # Bold "Setup." lead-in, as in "What we did:".
    r0 = setup.find(".//" + A + "r")
    lead = copy.deepcopy(r0)
    lead.find(A + "t").text = "Setup. "
    rpr = lead.find(A + "rPr")
    if rpr is None:
        rpr = etree.SubElement(lead, A + "rPr")
        lead.insert(0, rpr)
    rpr.set("b", "1")
    r0.addprevious(lead)

    prs.save(dst)
    print("wrote", dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
